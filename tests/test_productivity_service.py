"""Local productivity artifact tests."""

from datetime import UTC, datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from jarvis.application.productivity_service import LocalProductivityService
from jarvis.capabilities.language import LanguageMessage
from jarvis.capabilities.productivity import ProductivityError


class ContentModel:
    def complete(self, messages: list[LanguageMessage]) -> str:
        instruction = messages[0].content
        if "CSV only" in instruction:
            return "Item,Status;Jarvis,Ready"
        if "slide points" in instruction:
            return "Objective;Architecture;Validation"
        return "Executive summary\n\nGenerated from the supplied brief."


def test_email_draft_is_local_and_interoperable(tmp_path: Path) -> None:
    service = LocalProductivityService(tmp_path)
    result = service.create_email_draft("person@example.com", "Project update", "Done today.")

    message = BytesParser(policy=policy.default).parsebytes(result.path.read_bytes())
    assert result.kind == "email draft"
    assert message["To"] == "person@example.com"
    assert message["Subject"] == "Project update"
    assert "Done today." in message.get_content()


def test_calendar_event_is_utc_ics(tmp_path: Path) -> None:
    service = LocalProductivityService(tmp_path)
    result = service.create_calendar_event(
        "Jarvis review", datetime(2026, 8, 2, 12, 0, tzinfo=UTC), 45
    )
    contents = result.path.read_text(encoding="utf-8")

    assert result.path.suffix == ".ics"
    assert "DTSTART:20260802T120000Z" in contents
    assert "DTEND:20260802T124500Z" in contents
    assert "SUMMARY:Jarvis review" in contents


def test_documents_do_not_overwrite_existing_artifacts(tmp_path: Path) -> None:
    service = LocalProductivityService(tmp_path)
    first = service.create_document("report", "Weekly Status", "First version")
    second = service.create_document("report", "Weekly Status", "Second version")

    assert first.path != second.path
    assert "First version" in "\n".join(item.text for item in Document(first.path).paragraphs)
    assert "Second version" in "\n".join(item.text for item in Document(second.path).paragraphs)


def test_productivity_input_validation(tmp_path: Path) -> None:
    service = LocalProductivityService(tmp_path)
    with pytest.raises(ProductivityError, match="email"):
        service.create_email_draft("not-an-email", "Subject", "Body")
    with pytest.raises(ProductivityError, match="duration"):
        service.create_calendar_event("Meeting", datetime.now(UTC), 0)
    with pytest.raises(ProductivityError, match="type"):
        service.create_document("unknown", "Title", "Body")


def test_real_powerpoint_and_spreadsheet_are_generated(tmp_path: Path) -> None:
    service = LocalProductivityService(tmp_path)
    deck = service.create_presentation("Project Plan", "Objective; Milestone; Next action")
    sheet = service.create_spreadsheet("Budget", "Item,Cost;Model,523;Voice,0")

    presentation = Presentation(deck.path)
    workbook = load_workbook(sheet.path)

    assert deck.path.suffix == ".pptx" and len(presentation.slides) == 2
    assert sheet.path.suffix == ".xlsx"
    assert workbook.active["A2"].value == "Model"
    assert workbook.active["B2"].value == "523"


def test_model_generates_real_editable_artifacts_from_a_brief(tmp_path: Path) -> None:
    service = LocalProductivityService(tmp_path, ContentModel())
    document = service.generate_document("report", "AI Review", "Review local AI")
    presentation = service.generate_presentation("AI Slides", "Explain local AI")
    spreadsheet = service.generate_spreadsheet("AI Tracker", "Track readiness")

    assert Document(document.path).paragraphs[-1].text == "Generated from the supplied brief."
    assert len(Presentation(presentation.path).slides) == 2
    assert load_workbook(spreadsheet.path).active["A2"].value == "Jarvis"


def test_model_generation_requires_a_configured_local_model(tmp_path: Path) -> None:
    with pytest.raises(ProductivityError, match="unavailable"):
        LocalProductivityService(tmp_path).generate_document("report", "Title", "Brief")
